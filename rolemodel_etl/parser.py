from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import (
    AccessEntry,
    EntitlementRecord,
    JustificationEntry,
    ListEntry,
    ParseIssue,
    ParseResult,
    ProfileRecord,
    StructurePath,
    SystemRecord,
)


DEFAULT_SHEET_NAME = "Полная форма"
PROFILE_CODE_RE = re.compile(r"№:\s*(P\d+)", re.IGNORECASE)
PROFILE_TYPE_RE = re.compile(r"Тип профиля:\s*([^\)\n]+)", re.IGNORECASE)
MODEL_CODE_RE = re.compile(r"(RM\d+)", re.IGNORECASE)
CI_CODE_RE = re.compile(r"\[(CI\d+)\]", re.IGNORECASE)
LIST_LINE_RE = re.compile(r"^(.*?)\((\d{4,})\)\s*;?$")
ACCESS_VALUE_RE = re.compile(r"^\s*([12])(?:\s+(.*?))?\s*$", re.DOTALL)


def normalize_whitespace(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def stringify_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def parse_access_value(value: object) -> Optional[tuple[int, Optional[str], str]]:
    raw_text = stringify_cell(value)
    if not raw_text.strip():
        return None
    normalized = normalize_whitespace(raw_text)
    match = ACCESS_VALUE_RE.match(normalized)
    if not match:
        raise ValueError(f"Invalid access value: {normalized}")
    level = int(match.group(1))
    comment = normalize_whitespace(match.group(2)) if match.group(2) else None
    if comment == "":
        comment = None
    return level, comment, normalized


def parse_header_cell(value: object) -> tuple[str, str]:
    text = stringify_cell(value)
    parts = [line.strip() for line in text.splitlines() if line.strip()]
    if not parts:
        return "", ""
    first_line = parts[0]
    if len(parts) == 1 and ":" in first_line:
        head, tail = first_line.split(":", 1)
        header_type = normalize_whitespace(head)
        header_name = normalize_whitespace(tail)
    else:
        header_type = first_line.rstrip(":").strip()
        header_name = normalize_whitespace(" ".join(parts[1:])) if len(parts) > 1 else ""
    return header_type, header_name


def choose_sheet_name(sheet_names: list[str], requested_sheet: Optional[str]) -> str:
    if requested_sheet:
        if requested_sheet not in sheet_names:
            raise ValueError(
                f"Sheet '{requested_sheet}' not found. Available: {', '.join(sheet_names)}"
            )
        return requested_sheet
    if DEFAULT_SHEET_NAME in sheet_names:
        return DEFAULT_SHEET_NAME
    return sheet_names[0]


def _header_contains(header_value: object, required_fragments: list[str]) -> bool:
    header = normalize_whitespace(stringify_cell(header_value).lower())
    return all(fragment in header for fragment in required_fragments)


def _parse_list_entries(raw_text: str, heading_prefix: str) -> list[ListEntry]:
    entries: list[ListEntry] = []
    for line in [line.strip() for line in raw_text.splitlines() if line.strip()]:
        if line.lower().startswith(heading_prefix):
            continue
        match = LIST_LINE_RE.match(line)
        if match:
            name = normalize_whitespace(match.group(1))
            code = match.group(2)
            entries.append(
                ListEntry(name=name, code=code, raw_line=line, parse_status="PARSED")
            )
        else:
            entries.append(
                ListEntry(name=normalize_whitespace(line), code=None, raw_line=line, parse_status="RAW_ONLY")
            )
    return entries


def _make_issue(
    issues: list[ParseIssue],
    severity: str,
    stage: str,
    sheet_name: str,
    source_row: Optional[int],
    source_col: Optional[int],
    error_code: str,
    message: str,
    raw_value: Optional[str] = None,
) -> None:
    issues.append(
        ParseIssue(
            severity=severity,
            stage=stage,
            sheet_name=sheet_name,
            source_row=source_row,
            source_col=source_col,
            error_code=error_code,
            message=message,
            raw_value=raw_value,
        )
    )


def _validate_required_headers(ws: Worksheet, sheet_name: str, issues: list[ParseIssue]) -> None:
    required = {
        1: ["профиль"],
        2: ["обоснование"],
        3: ["структура"],
        4: ["подраздел", "команд", "атрибут"],
        5: ["должност", "рол"],
    }
    for col, fragments in required.items():
        if not _header_contains(ws.cell(1, col).value, fragments):
            _make_issue(
                issues=issues,
                severity="ERROR",
                stage="VALIDATE",
                sheet_name=sheet_name,
                source_row=1,
                source_col=col,
                error_code="HEADER_MISMATCH",
                message=f"Column {col} header does not match expected pattern: {fragments}",
                raw_value=stringify_cell(ws.cell(1, col).value),
            )


def parse_workbook(file_path: str | Path, requested_sheet: Optional[str] = None) -> ParseResult:
    path = Path(file_path)
    wb = load_workbook(path, data_only=True)
    sheet_name = choose_sheet_name(wb.sheetnames, requested_sheet)
    ws = wb[sheet_name]

    issues: list[ParseIssue] = []
    _validate_required_headers(ws, sheet_name, issues)

    model_meta_raw = normalize_whitespace(stringify_cell(ws.cell(1, 6).value))
    model_code_match = MODEL_CODE_RE.search(model_meta_raw)
    model_code = model_code_match.group(1).upper() if model_code_match else None
    model_name = (
        normalize_whitespace(MODEL_CODE_RE.sub("", model_meta_raw)).strip() if model_meta_raw else None
    )
    if model_name == "":
        model_name = None

    active_cols = []
    for col in range(6, ws.max_column + 1):
        if stringify_cell(ws.cell(3, col).value).strip():
            active_cols.append(col)

    if not active_cols:
        _make_issue(
            issues=issues,
            severity="ERROR",
            stage="VALIDATE",
            sheet_name=sheet_name,
            source_row=3,
            source_col=6,
            error_code="NO_ACTIVE_COLUMNS",
            message="No active entitlement columns found in row 3 from column 6 onwards.",
        )
        return ParseResult(
            source_file=str(path),
            sheet_name=sheet_name,
            model_code=model_code,
            model_name=model_name,
            rows_read=0,
            systems=[],
            entitlements=[],
            profiles=[],
            issues=issues,
        )

    system_by_col: dict[int, str] = {}
    current_system = ""
    for col in active_cols:
        row2_value = normalize_whitespace(stringify_cell(ws.cell(2, col).value))
        if row2_value:
            current_system = row2_value
        if not current_system:
            current_system = f"UNKNOWN_SYSTEM_COL_{col}"
            _make_issue(
                issues=issues,
                severity="ERROR",
                stage="VALIDATE",
                sheet_name=sheet_name,
                source_row=2,
                source_col=col,
                error_code="SYSTEM_HEADER_MISSING",
                message="Failed to resolve system name via forward-fill.",
            )
        system_by_col[col] = current_system

    reason_cols_by_system: dict[str, list[int]] = {}
    entitlements: list[EntitlementRecord] = []
    for col in active_cols:
        system_name = system_by_col[col]
        header_raw = stringify_cell(ws.cell(3, col).value)
        ent_type, ent_name = parse_header_cell(header_raw)
        if normalize_whitespace(ent_type.lower()) == "обоснование":
            reason_cols_by_system.setdefault(system_name, []).append(col)
            continue
        entitlements.append(
            EntitlementRecord(
                source_col=col,
                system_name_raw=system_name,
                entitlement_type=ent_type,
                entitlement_name=ent_name,
                header_raw=header_raw,
            )
        )

    systems_bounds: dict[str, tuple[int, int]] = {}
    for col in active_cols:
        sys_name = system_by_col[col]
        if sys_name not in systems_bounds:
            systems_bounds[sys_name] = (col, col)
        else:
            left, right = systems_bounds[sys_name]
            systems_bounds[sys_name] = (min(left, col), max(right, col))

    systems: list[SystemRecord] = []
    for system_name, (left, right) in systems_bounds.items():
        reason_cols = reason_cols_by_system.get(system_name, [])
        reason_col: Optional[int] = None
        if not reason_cols:
            _make_issue(
                issues=issues,
                severity="ERROR",
                stage="VALIDATE",
                sheet_name=sheet_name,
                source_row=3,
                source_col=left,
                error_code="REASON_COLUMN_MISSING",
                message=f"System '{system_name}' has no 'Обоснование' column.",
            )
        else:
            reason_col = reason_cols[0]
            if len(reason_cols) > 1:
                _make_issue(
                    issues=issues,
                    severity="WARNING",
                    stage="VALIDATE",
                    sheet_name=sheet_name,
                    source_row=3,
                    source_col=reason_cols[1],
                    error_code="MULTIPLE_REASON_COLUMNS",
                    message=f"System '{system_name}' has multiple 'Обоснование' columns; first one is used.",
                )

        ci_match = CI_CODE_RE.search(system_name)
        ci_code = ci_match.group(1).upper() if ci_match else None
        systems.append(
            SystemRecord(
                system_name_raw=system_name,
                ci_code=ci_code,
                source_col_start=left,
                source_col_end=right,
                reason_col=reason_col,
            )
        )

    systems_by_name = {system.system_name_raw: system for system in systems}

    profiles: list[ProfileRecord] = []
    rows_read = 0
    for row in range(4, ws.max_row + 1):
        raw_profile_cell = stringify_cell(ws.cell(row, 1).value)
        if not raw_profile_cell.strip():
            continue
        rows_read += 1

        profile_raw = raw_profile_cell
        profile_code_match = PROFILE_CODE_RE.search(profile_raw)
        if not profile_code_match:
            _make_issue(
                issues=issues,
                severity="ERROR",
                stage="PARSE_PROFILE",
                sheet_name=sheet_name,
                source_row=row,
                source_col=1,
                error_code="PROFILE_CODE_MISSING",
                message="Profile row skipped because profile code was not found.",
                raw_value=normalize_whitespace(profile_raw),
            )
            continue

        profile_code = profile_code_match.group(1).upper()
        profile_lines = [line.strip() for line in profile_raw.splitlines() if line.strip()]
        profile_name = profile_lines[0] if profile_lines else profile_code
        profile_type_match = PROFILE_TYPE_RE.search(profile_raw)
        profile_type = profile_type_match.group(1).strip() if profile_type_match else None
        profile_justification_raw = normalize_whitespace(stringify_cell(ws.cell(row, 2).value)) or None

        structure_paths: list[StructurePath] = []
        structure_raw = stringify_cell(ws.cell(row, 3).value)
        for path_idx, path_line in enumerate(
            [line.strip() for line in structure_raw.splitlines() if line.strip()], start=1
        ):
            segments = [segment.strip() for segment in path_line.split("\\") if segment.strip()]
            structure_paths.append(
                StructurePath(path_order=path_idx, path_raw=path_line, segments=segments)
            )

        departments = _parse_list_entries(
            stringify_cell(ws.cell(row, 4).value),
            heading_prefix="выбранные подразделения",
        )
        positions = _parse_list_entries(
            stringify_cell(ws.cell(row, 5).value),
            heading_prefix="выбранные должности",
        )

        profile = ProfileRecord(
            source_row=row,
            profile_code=profile_code,
            profile_name=profile_name,
            profile_type=profile_type,
            profile_raw=normalize_whitespace(profile_raw),
            profile_justification_raw=profile_justification_raw,
            structure_paths=structure_paths,
            departments=departments,
            positions=positions,
            justifications=[],
            accesses=[],
        )

        reason_map: dict[tuple[str, int], JustificationEntry] = {}
        for system in systems:
            if system.reason_col is None:
                continue
            reason_value = ws.cell(row, system.reason_col).value
            if not stringify_cell(reason_value).strip():
                continue
            try:
                parsed_reason = parse_access_value(reason_value)
            except ValueError:
                _make_issue(
                    issues=issues,
                    severity="ERROR",
                    stage="PARSE_JUSTIFICATION",
                    sheet_name=sheet_name,
                    source_row=row,
                    source_col=system.reason_col,
                    error_code="JUSTIFICATION_INVALID_VALUE",
                    message=f"Invalid justification value for system '{system.system_name_raw}'.",
                    raw_value=normalize_whitespace(stringify_cell(reason_value)),
                )
                continue
            if parsed_reason is None:
                continue
            level, comment, raw = parsed_reason
            justification = JustificationEntry(
                system_name_raw=system.system_name_raw,
                access_level=level,
                justification_text=comment,
                raw_value=raw,
                source_col=system.reason_col,
            )
            reason_key = (system.system_name_raw, level)
            reason_map[reason_key] = justification
            profile.justifications.append(justification)

        for entitlement in entitlements:
            cell_value = ws.cell(row, entitlement.source_col).value
            if not stringify_cell(cell_value).strip():
                continue
            try:
                parsed_access = parse_access_value(cell_value)
            except ValueError:
                _make_issue(
                    issues=issues,
                    severity="ERROR",
                    stage="PARSE_ACCESS",
                    sheet_name=sheet_name,
                    source_row=row,
                    source_col=entitlement.source_col,
                    error_code="ACCESS_INVALID_VALUE",
                    message=(
                        f"Invalid access value for entitlement '{entitlement.entitlement_type} / "
                        f"{entitlement.entitlement_name}'."
                    ),
                    raw_value=normalize_whitespace(stringify_cell(cell_value)),
                )
                continue
            if parsed_access is None:
                continue
            access_level, inline_comment, raw = parsed_access
            if (entitlement.system_name_raw, access_level) not in reason_map:
                _make_issue(
                    issues=issues,
                    severity="WARNING",
                    stage="PARSE_ACCESS",
                    sheet_name=sheet_name,
                    source_row=row,
                    source_col=entitlement.source_col,
                    error_code="JUSTIFICATION_NOT_FOUND",
                    message=(
                        f"Justification not found for system '{entitlement.system_name_raw}' "
                        f"and access level '{access_level}'."
                    ),
                    raw_value=raw,
                )
            profile.accesses.append(
                AccessEntry(
                    source_col=entitlement.source_col,
                    system_name_raw=entitlement.system_name_raw,
                    access_level=access_level,
                    inline_comment=inline_comment,
                    raw_value=raw,
                    source_row=row,
                )
            )

        profiles.append(profile)

    # Keep system order stable by source column.
    systems.sort(key=lambda item: item.source_col_start)
    entitlements.sort(key=lambda item: item.source_col)

    return ParseResult(
        source_file=str(path),
        sheet_name=sheet_name,
        model_code=model_code,
        model_name=model_name,
        rows_read=rows_read,
        systems=systems,
        entitlements=entitlements,
        profiles=profiles,
        issues=issues,
    )
